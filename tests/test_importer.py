import pandas as pd
import pytest
from docx import Document

from app.processing.importer import (ImportOptions, commit, held_in_database, import_file,
                                     load_records, plan_groups)
from app.processing.parsers import ParseError, read_file

from .conftest import numbers


@pytest.fixture
def csv_file(tmp_path):
    path = tmp_path / "with_copy.csv"
    path.write_text(
        "Bond Number,Denomination,Date of Purchase,Copy,Position\n"
        "100003,Rs. 750,2025-01-10,Alpha,3\n"
        "100001,750,10/01/2025,Alpha,1\n"
        "100002,750,2025-01-10,Alpha,2\n"
        "555010,1500,2025-03-01,Beta,\n"
        "555011,1500,2025-03-01,Beta,\n"
        "555010,1500,2025-03-01,Beta,\n"
        "ABC12,750,,Alpha,\n", encoding="utf-8")
    return path


@pytest.fixture
def xlsx_file(tmp_path):
    path = tmp_path / "two_sheets.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"Bond No": [123456, 123457, 4512], "Price": [100, 100, 100]}).to_excel(
            writer, sheet_name="Hundreds", index=False)
        pd.DataFrame({0: [900001, 900002, "900003-900005"]}).to_excel(
            writer, sheet_name="Misc", index=False, header=False)
    return path


@pytest.fixture
def docx_file(tmp_path):
    path = tmp_path / "list.docx"
    doc = Document()
    doc.add_paragraph("My prize bonds - 2026 (page 1 of 2)")
    doc.add_paragraph("Copy 1: 100001-100010")
    doc.add_paragraph("Copy 2")
    doc.add_paragraph("200001-200003, 200010")
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Bond Number", "Denomination"), ("300001", "200"), ("300002", "200")]):
        table.rows[r].cells[0].text, table.rows[r].cells[1].text = a, b
    doc.save(path)
    return path


def test_csv_columns_copy_column_positions_and_problems(csv_file):
    parsed = load_records(csv_file)
    assert len(parsed.records) == 6 and parsed.duplicates == []     # nothing is merged at file level
    assert any("ABC12" in p for p in parsed.problems)
    assert parsed.has_copy_column and parsed.has_position_column and not parsed.missing_denomination
    groups = {g.name: g for g in plan_groups(parsed, ImportOptions())}
    assert set(groups) == {"Alpha", "Beta"}
    assert groups["Alpha"].positions == [3, 1, 2] and groups["Beta"].positions is None
    assert groups["Beta"].repeated == ["555010"] and len(groups["Beta"].items) == 2   # once per copy
    assert {r.purchase_date for r in groups["Alpha"].items} == {"2025-01-10"}   # 10/01/2025 is day-first


def test_csv_import_honours_file_positions(db, csv_file):
    result = import_file(db, csv_file)
    assert result.added == 5 and result.copy_names == ["Alpha", "Beta"]
    alpha = next(c for c in db.list_copies() if c["name"] == "Alpha")
    assert alpha["first_number"] == "100001"
    bonds = db.copy_bonds(alpha["id"])
    assert [(b["number"], b["position"]) for b in bonds] == [
        ("100001", 1), ("100002", 2), ("100003", 3)]
    assert {b["denomination"] for b in bonds} == {750}


def test_excel_each_sheet_becomes_a_copy_and_leading_zeros_return(db, xlsx_file):
    parsed = load_records(xlsx_file)
    assert parsed.sources == ["Hundreds", "Misc"] and parsed.missing_denomination
    with pytest.raises(ValueError):                    # 'Misc' has no denomination
        commit(db, plan_groups(parsed, ImportOptions()), ImportOptions())
    assert db.list_bonds() == []                       # all-or-nothing
    options = ImportOptions(default_denomination=750, default_date="2026-01-01")
    result = import_file(db, xlsx_file, options)
    assert result.added == 8
    assert result.copy_names == ["two_sheets - Hundreds", "two_sheets - Misc"]
    assert "004512" in {b["number"] for b in db.list_bonds(denomination=100)}


def test_word_headings_name_the_copies_and_tables_are_read(db, docx_file):
    parsed = load_records(docx_file)
    assert parsed.has_copy_column and len(parsed.records) == 16
    result = import_file(db, docx_file, ImportOptions(default_denomination=750))
    assert result.copy_names == ["Copy 1", "Copy 2", "list - Table 1"]
    copy2 = next(c for c in db.list_copies() if c["name"] == "Copy 2")
    assert copy2["size"] == 4 and copy2["first_number"] == "200001"
    table_bonds = db.list_bonds(denomination=200)
    assert len(table_bonds) == 2                       # denomination column wins over the default


def test_split_into_chunks_and_loose_and_single(db, tmp_path):
    path = tmp_path / "plain.csv"
    path.write_text("\n".join(str(n) for n in range(100001, 100251)), encoding="utf-8")
    parsed = load_records(path)
    assert len(parsed.records) == 250 and not parsed.problems
    chunks = plan_groups(parsed, ImportOptions(mode="chunks", chunk_size=100, copy_name="Box"))
    assert [(g.name, len(g.items)) for g in chunks] == [("Box - 1", 100), ("Box - 2", 100), ("Box - 3", 50)]
    assert plan_groups(parsed, ImportOptions(mode="none"))[0].name is None
    assert plan_groups(parsed, ImportOptions(mode="single"))[0].name == "plain"
    result = import_file(db, path, ImportOptions(mode="chunks", chunk_size=100,
                                                 default_denomination=100))
    assert result.added == 250 and len(result.copy_names) == 3


def test_reimporting_the_same_file_adds_another_copy_with_the_same_numbers(db, tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("Bond Number\n100001\n100002\n", encoding="utf-8")
    options = ImportOptions(default_denomination=100)
    assert import_file(db, path, options).added == 2
    again = import_file(db, path, options)
    assert again.added == 2 and again.shared == 2 and not again.skipped_duplicates
    assert again.copy_names == ["a (2)"]                      # a second copy, same numbers
    assert {b["same_count"] for b in db.list_bonds()} == {2}


def test_skip_existing_leaves_out_numbers_already_in_the_database(db, tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("Bond Number\n100001\n100002\n", encoding="utf-8")
    options = ImportOptions(default_denomination=100, skip_existing=True)
    assert import_file(db, path, options).added == 2
    again = import_file(db, path, options)
    assert again.added == 0 and again.skipped_duplicates == ["100001", "100002"]
    assert len(db.list_copies()) == 1                         # no empty copy created


def test_copies_in_one_file_may_share_numbers_and_skip_only_looks_at_older_bonds(db, tmp_path):
    path = tmp_path / "books.csv"
    path.write_text("Copy,Bond Number\nA,100001\nA,100002\nA,100003\nB,100001\nB,100002\n"
                    "B,100003\nB,100003\n", encoding="utf-8")
    parsed = load_records(path)
    groups = {g.name: g for g in plan_groups(parsed, ImportOptions())}
    assert [len(g.items) for g in groups.values()] == [3, 3] and groups["B"].repeated == ["100003"]
    db.add_loose_bonds(numbers("100002"), 100)                # held before the import
    result = import_file(db, path, ImportOptions(default_denomination=100, skip_existing=True))
    assert result.added == 4 and result.skipped_duplicates == ["100002", "100002"]
    assert {b["same_count"] for b in db.list_bonds() if b["number"] == "100001"} == {2}
    plain = import_file(db, path, ImportOptions(default_denomination=100))
    assert plain.added == 6 and plain.copy_names == ["A (2)", "B (2)"]


def test_held_in_database_counts_numbers_already_on_file(db, tmp_path):
    path = tmp_path / "a.csv"
    path.write_text("Bond Number\n100001\n100002\n100003\n", encoding="utf-8")
    parsed = load_records(path)
    assert held_in_database(db, parsed.records, 100) == (0, [])
    db.add_loose_bonds(numbers("100002-100003"), 100)
    assert held_in_database(db, parsed.records, 100) == (2, ["100002", "100003"])
    assert held_in_database(db, parsed.records, 750) == (0, [])      # other price, other bond
    assert held_in_database(db, parsed.records, None) == (0, [])      # no price chosen yet


def test_unsupported_and_missing_files(tmp_path):
    with pytest.raises(ParseError):
        read_file(tmp_path / "nope.csv")
    bad = tmp_path / "x.pdf"
    bad.write_text("x")
    with pytest.raises(ParseError):
        read_file(bad)


def test_semicolon_csv_without_header(tmp_path):
    path = tmp_path / "semi.csv"
    path.write_text("1;100001;note\n2;100002;note\n", encoding="utf-8")
    parsed = load_records(path)
    assert [r.display for r in parsed.records] == ["100001", "100002"]
