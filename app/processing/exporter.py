"""Export tables to CSV or Excel (the format follows the file extension)."""
from __future__ import annotations

import re
from pathlib import Path

from .reports import win_table

Table = tuple[list[str], list[list]]     # (headers, rows)

EXPORT_DIALOG_FILTER = "Excel workbook (*.xlsx);;CSV file (*.csv)"


def _frame(headers, rows):
    import pandas as pd
    return pd.DataFrame([list(r) for r in rows], columns=list(headers))


def _sheet_name(name: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Sheet"
    candidate, n = clean, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = clean[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _style_sheet(sheet, headers) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", fgColor="E5E7EB")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        cells = list(sheet[get_column_letter(index)])[1:200]
        longest = max([len(str(header))] + [len(str(c.value)) for c in cells if c.value is not None])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)


def write_workbook(path: str | Path, sheets: dict[str, Table]) -> Path:
    """One .xlsx file with a formatted sheet per table."""
    import pandas as pd

    path = Path(path)
    used: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, (headers, rows) in sheets.items():
            title = _sheet_name(name, used)
            _frame(headers, rows).to_excel(writer, sheet_name=title, index=False)
            _style_sheet(writer.sheets[title], headers)
    return path


def write_table(path: str | Path, headers: list[str], rows: list[list],
                sheet_name: str = "Data") -> Path:
    """Write one table; ``.csv`` or ``.xlsx`` decided by the extension."""
    path = Path(path)
    extension = path.suffix.lower()
    if extension == ".csv":
        # utf-8 with BOM so Excel opens it with the right encoding
        _frame(headers, rows).to_csv(path, index=False, encoding="utf-8-sig")
    elif extension == ".xlsx":
        write_workbook(path, {sheet_name: (headers, rows)})
    else:
        raise ValueError("Choose a file name ending in .csv or .xlsx.")
    return path


def write_csv_folder(folder: str | Path, sheets: dict[str, Table]) -> list[Path]:
    """One CSV file per table inside ``folder``."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for name, (headers, rows) in sheets.items():
        target = folder / (re.sub(r"[^\w\- ]+", "", name).strip() + ".csv")
        written.append(write_table(target, headers, rows))
    return written


# ------------------------------------------------------- table builders
def bonds_table(bonds: list[dict]) -> Table:
    headers = ["Bond number", "Denomination", "Purchase date", "Status", "Win pending",
               "Bonds with this number", "Copy (book)", "Position in copy",
               "1st bond number of copy"]
    rows = [[b["number"], b["denomination"], b["purchase_date"] or "",
             "Active" if b["status"] == "active" else "Won",
             "Yes" if b.get("pending_count") else "",
             b.get("same_count") or 1,
             b["copy_name"] or "", b["position"], b["copy_first_number"] or ""]
            for b in bonds]
    return headers, rows


def copies_table(copies: list[dict]) -> Table:
    headers = ["Copy (book)", "1st bond number", "Size (positions)", "Bonds on file",
               "Active", "Won", "Denomination", "First purchase date", "Note"]

    def denomination(c):
        low, high = c["min_denomination"], c["max_denomination"]
        return "" if low is None else (low if low == high else f"{low}-{high}")

    rows = [[c["name"], c["first_number"] or "", c["size"], c["bonds"], c["active"], c["won"],
             denomination(c), c["purchase_date"] or "", c["note"] or ""] for c in copies]
    return headers, rows


def draws_table(draws: list[dict]) -> Table:
    headers = ["Draw date", "Denomination", "Label", "Winning numbers", "Pending", "Approved wins"]
    rows = [[d["draw_date"], d["denomination"] if d["denomination"] else "All",
             d["label"] or "", d["numbers"], d["pending"], d["wins"]] for d in draws]
    return headers, rows


def winning_numbers_table(db) -> Table:
    headers = ["Draw date", "Draw", "Denomination", "Winning number", "Prize", "Prize amount"]
    data = db.query(
        "SELECT d.draw_date, d.label, d.denomination, w.number, w.prize_label, w.prize_amount "
        "FROM winning_numbers w JOIN draws d ON d.id = w.draw_id "
        "ORDER BY d.draw_date DESC, d.id DESC, w.prize_label, w.num")
    rows = [[r["draw_date"], r["label"] or "", r["denomination"] or "All", r["number"],
             r["prize_label"], r["prize_amount"]] for r in data]
    return headers, rows


def activity_table(entries: list[dict]) -> Table:
    return ["When", "Action", "Details"], [[e["ts"], e["action"], e["details"] or ""]
                                           for e in entries]


def collect_all(db) -> dict[str, Table]:
    """Every list in the app, ready for ``write_workbook`` / ``write_csv_folder``."""
    sheets: dict[str, Table] = {"All bonds": bonds_table(db.list_bonds())}
    for denomination in db.denominations("active"):
        sheets[f"Waiting Rs {denomination}"] = bonds_table(
            db.list_bonds(status="active", denomination=denomination))
    sheets["Copies"] = copies_table(db.list_copies())
    sheets["Pending wins"] = win_table(db.list_pending(), "Found on", "found_at")
    sheets["Win list"] = win_table(db.list_wins())
    sheets["Draws"] = draws_table(db.list_draws())
    sheets["Winning numbers"] = winning_numbers_table(db)
    sheets["Activity log"] = activity_table(db.list_activity())
    return sheets
