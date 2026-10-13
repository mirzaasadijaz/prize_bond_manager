"""Read CSV, Excel, Word and text files into plain tables of text cells.

This module only reads files.  Turning cells into bonds (headers, ranges,
copy names, denominations, dates) happens in ``importer.py``.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".csv": "CSV", ".tsv": "CSV", ".xlsx": "Excel", ".xlsm": "Excel",
    ".docx": "Word", ".txt": "Text",
}
FILE_DIALOG_FILTER = (
    "Bond lists (*.csv *.tsv *.xlsx *.xlsm *.docx *.txt);;"
    "CSV (*.csv *.tsv);;Excel (*.xlsx *.xlsm);;Word (*.docx);;Text (*.txt);;All files (*)"
)


class ParseError(ValueError):
    """A file could not be read; the message is safe to show to the user."""


@dataclass
class RawTable:
    """One sheet, Word table, CSV body or block of free text."""

    name: str
    rows: list[list[str]] = field(default_factory=list)
    kind: str = "table"            # "table" (cells) or "text" (one line per row)


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _read_csv(path: Path) -> list[RawTable]:
    text = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        cells = [c.strip() for c in row]
        if any(cells):
            rows.append(cells)
    return [RawTable(path.stem, rows)] if rows else []


def _read_text_file(path: Path) -> list[RawTable]:
    lines = [ln.strip() for ln in _read_text(path).splitlines() if ln.strip()]
    return [RawTable("Text", [[ln] for ln in lines], kind="text")] if lines else []


def _read_excel(path: Path) -> list[RawTable]:
    import pandas as pd

    try:
        sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=str,
                               keep_default_na=False, engine="openpyxl")
    except Exception as exc:  # corrupt file, password protected, wrong format ...
        raise ParseError(f"Could not read the Excel file: {exc}") from exc
    tables = []
    for name, frame in sheets.items():
        rows = [[str(c).strip() for c in row] for row in frame.itertuples(index=False, name=None)]
        rows = [r for r in rows if any(r)]
        if rows:
            tables.append(RawTable(str(name), rows))
    return tables


def _read_docx(path: Path) -> list[RawTable]:
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ParseError(f"Could not read the Word document: {exc}") from exc
    tables = []
    lines = [ln.strip() for p in document.paragraphs for ln in p.text.splitlines() if ln.strip()]
    if lines:
        tables.append(RawTable("Text", [[ln] for ln in lines], kind="text"))
    for index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(RawTable(f"Table {index}", rows))
    return tables


_READERS = {"CSV": _read_csv, "Excel": _read_excel, "Word": _read_docx, "Text": _read_text_file}


def read_file(path: str | Path) -> list[RawTable]:
    """Read any supported file into raw tables (empty list if it has no content)."""
    path = Path(path)
    kind = SUPPORTED_EXTENSIONS.get(path.suffix.lower())
    if kind is None:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ParseError(f"'{path.suffix or path.name}' files are not supported "
                         f"(supported: {supported}).")
    if not path.is_file():
        raise ParseError(f"File not found: {path}")
    return _READERS[kind](path)
